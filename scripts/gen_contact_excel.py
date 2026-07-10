#!/usr/bin/env python3
"""将 contact.yaml 测试用例生成 Excel 文件"""

import yaml
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

yaml_path = os.path.join(
    os.path.dirname(__file__), "..",
    "doc/SM测试用例-收集箱/AI输出/基础功能用例/contact.yaml"
)
yaml_path = os.path.abspath(yaml_path)

output_path = os.path.join(
    os.path.dirname(__file__), "..",
    "doc/SM测试用例-收集箱/AI输出/基础功能用例/contact.xlsx"
)
output_path = os.path.abspath(output_path)

with open(yaml_path, "r", encoding="utf-8") as f:
    data = yaml.safe_load(f)

cases = data.get("cases", [])

# --- Workbook ---
wb = Workbook()
ws = wb.active
ws.title = "联系人测试用例"

# --- Headers ---
headers = ["用例ID", "模块", "用例名称", "分类", "角色", "优先级",
           "前置条件", "测试步骤", "预期结果", "测试状态", "备注"]

# --- Styles ---
header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

cell_font = Font(name="微软雅黑", size=10)
cell_align = Alignment(vertical="top", wrap_text=True)
cell_align_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Write headers
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write data
def flatten(text):
    """将多行文本转为 Excel 可读的换行文本"""
    if text is None:
        return ""
    if isinstance(text, list):
        return "\n".join(str(item).strip() for item in text)
    return str(text).strip()

priority_colors = {
    "P0": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),  # 红
    "P1": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),  # 黄
    "P2": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),  # 绿
}

for row_idx, case in enumerate(cases, 2):
    values = [
        case.get("id", ""),
        case.get("module", ""),
        case.get("name", ""),
        case.get("category", ""),
        case.get("role", ""),
        case.get("priority", ""),
        flatten(case.get("precondition", "")),
        flatten(case.get("steps", "")),
        flatten(case.get("expected", "")),
        case.get("status", ""),
        flatten(case.get("notes", "")),
    ]

    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = cell_font
        cell.border = thin_border
        if col_idx in (1, 2, 4, 5, 6, 10):  # 短文本列居中
            cell.alignment = cell_align_center
        else:
            cell.alignment = cell_align

    # 优先级变色
    priority = case.get("priority", "")
    if priority in priority_colors:
        priority_cell = ws.cell(row=row_idx, column=6)
        priority_cell.fill = priority_colors[priority]

# --- Column widths ---
col_widths = {
    1: 18,   # 用例ID
    2: 10,   # 模块
    3: 40,   # 用例名称
    4: 16,   # 分类
    5: 10,   # 角色
    6: 8,    # 优先级
    7: 42,   # 前置条件
    8: 55,   # 测试步骤
    9: 50,   # 预期结果
    10: 10,  # 测试状态
    11: 50,  # 备注
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# 冻结首行
ws.freeze_panes = "A2"

# 自动筛选
ws.auto_filter.ref = f"A1:K{len(cases) + 1}"

# --- Save ---
wb.save(output_path)
print(f"Excel 已生成: {output_path}")
print(f"共 {len(cases)} 条用例")
