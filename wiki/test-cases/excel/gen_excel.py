#!/usr/bin/env python3
"""Generate English Excel test case file from CRM test cases markdown."""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ── colours ──────────────────────────────────────────────────────────────────
HEADER_BG   = "1F3864"   # dark navy
HEADER_FG   = "FFFFFF"
P0_BG       = "FF0000"
P1_BG       = "FF6600"
P2_BG       = "FFC000"
P3_BG       = "92D050"
SECTION_BG  = "D9E1F2"   # light blue-grey for alternate rows
ALT_BG      = "F2F7FF"

def pri_colour(pri):
    return {
        "P0": ("FF0000", "FFFFFF"),
        "P1": ("FF6600", "FFFFFF"),
        "P2": ("FFC000", "000000"),
        "P3": ("92D050", "000000"),
    }.get(pri, ("FFFFFF", "000000"))

def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

# ── data ─────────────────────────────────────────────────────────────────────
TEST_CASES = [
    {
        "id": "TC-001",
        "title": "Lead Conversion - Service Team and Activity Carry-Over Validation",
        "module": "Lead",
        "category": "Data Inheritance",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Lead exists with 2 Service Team members (Owner + Member) and 3 Activity Logs recorded",
        "steps": (
            "1. Navigate to Lead Details page\n"
            "2. Verify Service Team has 2 members\n"
            "3. Verify Activity Log has 3 entries\n"
            "4. Click 'Convert' to start 2-step conversion wizard\n"
            "5. Step 1 (Customer): Verify/Create customer, proceed\n"
            "6. Step 2 (Opportunity): Verify ALL 6 carry-over checkboxes are checked by default\n"
            "   - Copy Team to: Customer, Contact, Opportunity\n"
            "   - Copy Activities to: Customer, Contact, Opportunity\n"
            "7. Uncheck 'Copy Activities to -> Opportunity' checkbox\n"
            "8. Complete conversion\n"
            "9. Navigate to newly created Customer -> check Service Team\n"
            "10. Navigate to newly created Contact -> check Sales Record tab\n"
            "11. Navigate to newly created Opportunity -> check Service Team AND Activity Log"
        ),
        "expected": (
            "1. Customer Service Team contains both original Lead team members\n"
            "2. Contact Sales Record contains all 3 activity logs from Lead\n"
            "3. Opportunity Service Team contains both original Lead team members\n"
            "4. BUG REGRESSION: Opportunity Activity Log should NOT contain Lead activity logs (checkbox was unchecked)\n"
            "5. Lead status changes to 'Converted' (green badge)\n"
            "6. Conversion Record tab shows conversion history with timestamp"
        ),
    },
    {
        "id": "TC-002",
        "title": "Lead Conversion - Multi-Currency Amount Calculation (IDR to MYR)",
        "module": "Lead",
        "category": "Data Inheritance",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Lead exists with Currency=IDR, Estimated Deal Amount=100,000,000, Win Rate=25%",
        "steps": (
            "1. Navigate to Lead Details\n"
            "2. Record Estimated Deal Amount (100,000,000 IDR) and Win Rate (25%)\n"
            "3. Click 'Convert' -> Step 2 (Opportunity)\n"
            "4. Verify Currency defaults to IDR\n"
            "5. Enter Exchange Rate: 0.000223 (IDR->MYR)\n"
            "6. Verify MYR Value auto-calculates: 100,000,000 x 0.000223 = 22,300 MYR\n"
            "7. Verify Weighted Amount = MYR Value x Win Rate = 22,300 x 25% = 5,575 MYR\n"
            "8. Complete conversion\n"
            "9. Navigate to created Opportunity -> verify financial fields"
        ),
        "expected": (
            "1. MYR Value = 22,300.00 (exact: 100,000,000 x 0.000223)\n"
            "2. Weighted Amount = 5,575.00 (22,300 x 0.25)\n"
            "3. Opportunity Currency = IDR\n"
            "4. Entity reflects selected entity (SCMY or PTSM)\n"
            "5. Sales Record field auto-populated from Lead activity history\n"
            "6. Contacts tag input auto-carried from Lead Contact Person table"
        ),
    },
    {
        "id": "TC-003",
        "title": "Lead Queue - 7-Day SLA Reclaim Countdown",
        "module": "Lead",
        "category": "SLA Enforcement",
        "role": "Sales Manager",
        "priority": "P2",
        "precondition": "Lead exists in Unassigned queue, assigned 5 days ago, no follow-up activity",
        "steps": (
            "1. Login as Sales Manager\n"
            "2. Navigate to Lead Queue configuration\n"
            "3. Verify Claim and Assign Rule = 'Hidden from Members, Assignable by Admins'\n"
            "4. Navigate to Lead List -> filter by 'Unassigned' status\n"
            "5. Locate the lead assigned 5 days ago\n"
            "6. Verify 'Reclaim Countdown' column shows '2Days' (7 - 5 = 2 days remaining)\n"
            "7. Wait/simulate until countdown reaches '-1Days' (overdue)\n"
            "8. Verify lead is auto-returned to Public Pool\n"
            "9. Verify notification sent to original assignee"
        ),
        "expected": (
            "1. Reclaim Countdown accurately reflects remaining days\n"
            "2. Negative values (e.g., '-1Days') indicate overdue leads\n"
            "3. Overdue lead returns to Public Pool automatically\n"
            "4. Original assignee receives 'recycle countdown' notification\n"
            "5. Lead does NOT auto-strip - only returns to Public Pool (Blueprint principle 2)\n"
            "6. Any AM in the region can now claim the lead from Public Pool"
        ),
    },
    {
        "id": "TC-004",
        "title": "Lead Import - Enum Validation and Duplicate Check",
        "module": "Lead",
        "category": "Import Validation",
        "role": "Sales Manager",
        "priority": "P2",
        "precondition": "Excel template with 17 columns prepared (9 required, 8 optional)",
        "steps": (
            "1. Navigate to Lead List -> Click Import button\n"
            "2. Select Import Mode = 'New and update import' (UPSERT)\n"
            "3. Set Condition Field = 'Lead Name'\n"
            "4. Upload Excel template with test data:\n"
            "   - Row 1: Valid data, Lead Level='A-Level', Source='Search Engine'\n"
            "   - Row 2: Invalid Lead Level='D-Level' (not in enum)\n"
            "   - Row 3: Duplicate Lead Name matching existing record\n"
            "   - Row 4: Lock Status='1' (numeric -> should map to 'Unlock')\n"
            "5. Click Confirm Import\n"
            "6. Review import results"
        ),
        "expected": (
            "1. Row 1 imported successfully\n"
            "2. Row 2 fails validation - Lead Level 'D-Level' not in allowed values (A/B/C-Level)\n"
            "3. Row 3 updates existing record (UPSERT mode)\n"
            "4. Row 4 Lock Status correctly maps: 1=Unlock, 2=Locked\n"
            "5. Import summary shows: N succeeded, M failed, K updated\n"
            "6. All imported data visible in Lead Queue list view\n"
            "7. Enum: Lead Level (3 values), Source (7 values), Lead Queue (2 values)"
        ),
    },
    {
        "id": "TC-005",
        "title": "Lead Visibility - Entity Isolation and Service Team Access",
        "module": "Lead",
        "category": "Permission",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Lead-A (Entity=SCMY, User-A in Service Team), Lead-B (Entity=SMMY, User-B NOT in Service Team)",
        "steps": (
            "1. Login as User-A (SCMY entity)\n"
            "2. Navigate to Lead List -> filter 'All'\n"
            "3. Verify Lead-A is visible and editable\n"
            "4. Verify Lead-B is NOT visible (entity isolation)\n"
            "5. Login as User-B (SMMY entity)\n"
            "6. Navigate to Lead List\n"
            "7. Attempt to access Lead-A via direct URL\n"
            "8. Login as Admin, add User-B to Lead-A Service Team with Read-Only\n"
            "9. Login as User-B, attempt to edit Lead-A fields"
        ),
        "expected": (
            "1. Entity isolation: User-A sees only SCMY leads, User-B sees only SMMY leads\n"
            "2. Direct URL access to Lead-A from User-B returns access denied\n"
            "3. After adding to Service Team with Read-Only: User-B can view but CANNOT edit Lead-A\n"
            "4. Service Team permission overrides entity isolation for specific records\n"
            "5. Owner auto-assignment: Creator is added to Service Team automatically"
        ),
    },
    {
        "id": "TC-006",
        "title": "Opportunity Stage Progression - 5-Stage Pipeline with Win Rate",
        "module": "Opportunity",
        "category": "Pipeline Flow",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Opportunity exists at Qualifying stage (25% win rate)",
        "steps": (
            "1. Navigate to Opportunity Details\n"
            "2. Verify current stage = 'Qualifying' with 25% win rate badge\n"
            "3. Click 'Advance to the next stage'\n"
            "4. Verify stage changes to 'Proposal/POC' (50%)\n"
            "5. Verify Weighted Amount recalculates: Estimated Amount x 50%\n"
            "6. Click 'Advance to the next stage' again\n"
            "7. Verify stage changes to 'Price Negotiation' (75%)\n"
            "8. Click 'Advance to the next stage' -> Select 'Deal Won'\n"
            "9. Verify stage changes to 'Deal Won' (100%)\n"
            "10. Verify Status badge changes to 'Won' (green)\n"
            "11. Attempt to advance from 'Deal Won' -> should be blocked"
        ),
        "expected": (
            "1. Stage progression: Qualifying(25%) -> Proposal/POC(50%) -> Price Negotiation(75%) -> Deal Won(100%)\n"
            "2. Each stage advance recalculates Weighted Amount = Estimated Amount x Win Rate\n"
            "3. At Deal Won: Status = 'Won', Pipeline stage = 100%\n"
            "4. Deal Won is an End stage - cannot advance further\n"
            "5. Deal Lost (0%) is also an End stage\n"
            "6. All stage advances are manual (no auto-advance per Blueprint)"
        ),
    },
    {
        "id": "TC-007",
        "title": "Permission - Owner Without Service Team Access (Known Bug Regression)",
        "module": "Opportunity",
        "category": "Permission Bug",
        "role": "Sales Manager",
        "priority": "P1",
        "precondition": "Opportunity exists where Owner is NOT listed in Service Team members",
        "steps": (
            "1. Navigate to Opportunity Details where Owner is NOT any Service Team member\n"
            "2. Verify Owner field shows a user (e.g., 'Nixon')\n"
            "3. Open Service Team sidebar -> verify Owner is NOT listed\n"
            "4. As the Owner, attempt to create a Sales Order under this Opportunity\n"
            "5. As a non-owner NOT in Service Team, attempt to edit Opportunity fields\n"
            "6. As a non-owner NOT in Service Team, attempt to create a Sales Order"
        ),
        "expected": (
            "1. BUG: Owner CAN create Sales Order even without being in Service Team\n"
            "   - Expected: Owner should be automatically in Service Team\n"
            "   - Actual: Owner bypasses Service Team check for child record creation\n"
            "2. Non-owner NOT in Service Team CANNOT edit Opportunity fields (correct)\n"
            "3. Non-owner NOT in Service Team CANNOT create Sales Order (correct)\n"
            "4. Bug severity: HIGH - permission inconsistency"
        ),
    },
    {
        "id": "TC-008",
        "title": "Service Team - Add Member with Permission and Role Configuration",
        "module": "Opportunity",
        "category": "Service Team",
        "role": "Sales Manager",
        "priority": "P1",
        "precondition": "Opportunity exists with Owner only in Service Team",
        "steps": (
            "1. Navigate to Opportunity Details -> Service Team sidebar\n"
            "2. Click '+ Add' -> 'Add Team Members' modal opens\n"
            "3. Select member 'Solution Architect'\n"
            "4. Set Permission = 'Read-Write'\n"
            "5. Check Team Role = 'Ordinary Members' (default)\n"
            "6. Check Project Role = 'Customer Manager'\n"
            "7. Click Confirm\n"
            "8. Verify member appears in Service Team with correct roles\n"
            "9. Login as Solution Architect -> attempt to edit Opportunity fields\n"
            "10. Change Permission to 'Read-Only' for Solution Architect\n"
            "11. Login as Solution Architect -> attempt to edit -> should be blocked"
        ),
        "expected": (
            "1. New member appears in Service Team with all configured attributes\n"
            "2. Permission controls edit access: Read-Write = can edit, Read-Only = view only\n"
            "3. Team Role = 'Ordinary Members' defines internal team structure\n"
            "4. Project Role = 'Customer Manager' defines functional role on account\n"
            "5. Permission changes take effect immediately\n"
            "6. Permission is entity-specific: Opportunity Service Team != Customer/Lead access"
        ),
    },
    {
        "id": "TC-009",
        "title": "Activity Log - Interaction Recording Across Entity Types",
        "module": "Opportunity",
        "category": "Activity Log",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Opportunity exists with at least one Contact associated",
        "steps": (
            "1. Navigate to Opportunity Details -> Activity Log (right sidebar)\n"
            "2. Click 'New Log' / 'Log Activity'\n"
            "3. Fill Interaction Log form:\n"
            "   - Interaction Date: pre-filled with current datetime\n"
            "   - Contact: select from dropdown\n"
            "   - Our Attendee: pre-filled with Opportunity Owner\n"
            "   - Related Business: auto-filled with Opportunity Name\n"
            "   - Discussion Details: enter test text (verify 0/1000 char counter)\n"
            "   - Attachment: upload a test file\n"
            "4. Click Submit\n"
            "5. Verify log appears in Activity Log timeline\n"
            "6. Navigate to Contact Details -> check 'Sales Record' tab\n"
            "7. Verify same interaction appears"
        ),
        "expected": (
            "1. Activity log entry appears with: Date, User avatar, Activity Type badge, Timestamp, Content\n"
            "2. Metadata shows: 'Contact: [Name]', 'Our Attendee: [Owner]'\n"
            "3. BUG: Activity log tab naming inconsistent - 'Activity Log' vs 'Sales Record'\n"
            "4. Same activity should appear in both Opportunity Activity Log and Contact Sales Record\n"
            "5. Character counter shows current/max (e.g., '0/1000')\n"
            "6. Attachment is downloadable from activity log entry"
        ),
    },
    {
        "id": "TC-010",
        "title": "Opportunity Ownership Transfer - Team Continuity and Permission",
        "module": "Opportunity",
        "category": "Ownership Transfer",
        "role": "Sales Manager",
        "priority": "P2",
        "precondition": "Opportunity exists with Owner 'Affendi' and 2 Service Team members",
        "steps": (
            "1. Navigate to Opportunity Details -> Click 'Change Owner'\n"
            "2. Select new Owner = 'YCK'\n"
            "3. Set Original Owner disposition = 'Demote to team member'\n"
            "4. Set Permission = 'Read-Write' for original owner\n"
            "5. Set Team continuity = 'Keep Existing'\n"
            "6. Confirm transfer\n"
            "7. Verify new Owner = 'YCK' in header\n"
            "8. Verify original owner 'Affendi' still in Service Team with Read-Write\n"
            "9. Verify all existing Service Team members retained\n"
            "10. Login as original owner -> verify can still edit (Read-Write)\n"
            "11. Test 'Remove from team' option"
        ),
        "expected": (
            "1. New owner appears in Opportunity header and Service Team\n"
            "2. Original owner disposition applied correctly:\n"
            "   - 'Remove' -> removed from Service Team entirely\n"
            "   - 'Demote' -> stays in Service Team with specified permission\n"
            "3. Team continuity 'Keep Existing' retains all current members\n"
            "4. Team continuity 'Clear All' removes all existing members except new owner\n"
            "5. Notification sent to new owner about ownership transfer"
        ),
    },
    {
        "id": "TC-011",
        "title": "P&L Financial Summary - Margin Calculation Accuracy (Markup vs Gross Margin)",
        "module": "P&L",
        "category": "Financial Calculation",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "P&L exists with: 2 Software products (List Price: 5000, 3000) and 1 Hardware (Cost: 2000)",
        "steps": (
            "1. Navigate to P&L Create/Details page\n"
            "2. Add 2 Software products:\n"
            "   - Product A: List Price = 5,000, Markup = 20% -> Selling Price = 6,000\n"
            "   - Product B: List Price = 3,000, Markup = 10% -> Selling Price = 3,300\n"
            "3. Add 1 Hardware product:\n"
            "   - Cost per unit = 2,000, Expected Profit = 500 -> Selling Price = 2,500\n"
            "4. Verify Financial Summary Cards:\n"
            "   - TOTAL REVENUE = 6,000 + 3,300 + 2,500 = 11,800\n"
            "   - TOTAL COST = software cost + 2,000\n"
            "   - TOTAL PROFIT = TOTAL REVENUE - TOTAL COST\n"
            "   - MARGIN % = (TOTAL PROFIT / TOTAL REVENUE) x 100\n"
            "5. Verify each line item Margin = (Profit / Selling Price) x 100\n"
            "6. Verify red warning icon when Margin < Target Margin"
        ),
        "expected": (
            "1. CRITICAL: System uses Gross Margin formula (Profit/Selling Price)x100\n"
            "2. Software: Selling Price = List Price x (1 + Markup%)\n"
            "3. Hardware: Selling Price = Cost + Expected Profit\n"
            "4. TOTAL REVENUE card: blue text, sum of all selling prices\n"
            "5. TOTAL PROFIT card: green text (Revenue - Cost)\n"
            "6. MARGIN card: green if >= target, red if < target\n"
            "7. Red exclamation warning icon on line items where Margin < Target\n"
            "8. All figures auto-calculate"
        ),
    },
    {
        "id": "TC-012",
        "title": "P&L Global Discount - Toggle ON and OFF Behavior",
        "module": "P&L",
        "category": "Discount Control",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "P&L exists with 3 product line items (2 Software, 1 Hardware)",
        "steps": (
            "1. Navigate to P&L Details page\n"
            "2. Locate Global Discount bar (yellow-highlighted)\n"
            "3. Toggle ON, set discount = 5%\n"
            "4. Verify ALL line items show 5% discount in DISC column\n"
            "5. Verify PRICE column recalculates: PRICE = (List Price x (1+Markup%)) x (1-5%)\n"
            "6. Verify Financial Summary Cards update with new discounted values\n"
            "7. Toggle OFF\n"
            "8. Verify individual line item discount fields appear\n"
            "9. Set Product A discount = 10%, Product B = 0%, Hardware = 3%\n"
            "10. Verify each line item uses its own discount rate\n"
            "11. Verify Financial Summary reflects mixed discount rates"
        ),
        "expected": (
            "1. Toggle ON: Global discount applies uniformly to ALL line items\n"
            "2. Toggle OFF: Per-item discount fields appear for granular control\n"
            "3. Discount calculation: Net Price = Base Price x (1 - Discount%)\n"
            "4. Financial Summary always reflects current discount state\n"
            "5. Switching toggle preserves individual discount values\n"
            "6. Global discount toggle is blue when ON"
        ),
    },
    {
        "id": "TC-013",
        "title": "P&L Multi-User Collaboration - Sales Rep to Solution Architect Workflow",
        "module": "P&L",
        "category": "Collaboration",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "P&L exists in draft state, Solution Architect user exists",
        "steps": (
            "1. Login as Sales Rep (initiator)\n"
            "2. Navigate to P&L Details -> Process Approvals sidebar\n"
            "3. Set Operation = 'Collaborate' (NOT 'submit')\n"
            "4. Select Recipient = 'Solution Architect' via user picker\n"
            "5. Add collaboration comment: 'Please review hardware pricing'\n"
            "6. Click Submit -> form permission transfers to Recipient\n"
            "7. Login as Solution Architect -> check Message Center -> To-do Items\n"
            "8. Open P&L from To-do -> click Edit (pencil icon)\n"
            "9. Modify hardware cost, add professional services line\n"
            "10. Click 'Submit back' -> returns to Sales Rep\n"
            "11. Login as Sales Rep -> verify changes are visible\n"
            "12. Sales Rep can assign another collaborator OR click 'submit' for approval"
        ),
        "expected": (
            "1. Collaboration transfers edit permission from initiator to recipient\n"
            "2. Recipient receives To-do item in Message Center\n"
            "3. Recipient can edit ALL fields and add/remove line items\n"
            "4. 'Submit back' returns permission to initiator\n"
            "5. Initiator can repeat cycle with different collaborators\n"
            "6. Only after collaboration complete can initiator click 'submit' for approval\n"
            "7. Collaboration is distinct from approval - no HOD/Finance routing during collaboration"
        ),
    },
    {
        "id": "TC-014",
        "title": "P&L Approval Workflow - Margin-Based Routing Conditions",
        "module": "P&L",
        "category": "Approval Routing",
        "role": "Sales Manager",
        "priority": "P1",
        "precondition": "P&L with Product A: Margin=15% < Target 20%, Product B: Margin=25% >= Target 20%",
        "steps": (
            "1. Navigate to P&L Details\n"
            "2. Review line items:\n"
            "   - Product A: Margin = 15%, Target = 20% -> BELOW target\n"
            "   - Product B: Margin = 25%, Target = 20% -> ABOVE target\n"
            "3. Click 'submit' in Process Approvals sidebar\n"
            "4. Track approval route\n"
            "5. Verify route goes to Head of Department (HOD) due to Product A margin < target\n"
            "6. Test scenario: ALL products margin >= target -> verify auto-approve\n"
            "7. Test scenario: Grand total margin < target -> verify routes to Finance team"
        ),
        "expected": (
            "1. Approval routing rules:\n"
            "   - ANY line item margin < target -> route to HOD\n"
            "   - ANY Key Product included -> route to HOD\n"
            "   - Grand total margin < target -> route to Finance team\n"
            "   - All criteria met -> auto-approve\n"
            "2. Routes to HOD because Product A margin (15%) < target (20%)\n"
            "3. HOD can approve, reject, or request revision\n"
            "4. Rejected P&L must be revised and resubmitted\n"
            "5. Approved P&L -> Quote can be generated via 'New Quotation' button"
        ),
    },
    {
        "id": "TC-015",
        "title": "P&L Product Selection - Modal Filters and Category Scope",
        "module": "P&L",
        "category": "Product Selection",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "P&L exists, Opportunity has 3 products (2 Hardware, 1 Software)",
        "steps": (
            "1. Navigate to P&L Create -> click '+ Add Products' in SM Products section\n"
            "2. Verify Select Product Modal opens with dual-panel layout\n"
            "3. Left panel: Search filters (Product Code, Product Description)\n"
            "4. Right panel: 'Selected (0)' - empty initially\n"
            "5. Search for a product by code\n"
            "6. Select 2 products -> verify 'Selected (2)' updates\n"
            "7. Click 'Add 2 Products'\n"
            "8. Verify products appear in table with auto-populated: Code, Unit Price, Target Margin\n"
            "9. Navigate to Hardware section -> click '+ Add Products'\n"
            "10. BUG TEST: Search for Software product in Hardware section -> should NOT appear"
        ),
        "expected": (
            "1. Select Product Modal: dual-panel with search filters and selection counter\n"
            "2. Products auto-carry from Product Master: List Price, Cost, Target Margin\n"
            "3. BUG #1: Software products not selectable in P&L Software section - only Principal=Securemetric shows\n"
            "4. BUG #2: Opportunity products do not auto-fill in P&L product line\n"
            "5. BUG #5: Third-party product filter not scoped to entity/opportunity - shows ALL products\n"
            "6. Services section uses Service Master for role-based pricing\n"
            "7. Each section has separate '+ Add Products' button"
        ),
    },
    {
        "id": "TC-016",
        "title": "P&L Details - Multi-Year Breakdown (Year 1-3 RENEW)",
        "module": "P&L",
        "category": "Multi-Year",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Approved P&L exists with Hardware products spanning multiple years",
        "steps": (
            "1. Navigate to P&L Details page (not Create)\n"
            "2. Locate multi-year breakdown table\n"
            "3. Verify columns: CATEGORY | YEAR 1 | YEAR 2 (RENEW) | YEAR 3 (RENEW) | TOTAL\n"
            "4. Verify rows: Software, Hardware, Services, Reimbursement\n"
            "5. Verify each cell shows: Revenue, Cost, Margin % (color-coded)\n"
            "6. Verify GRAND TOTAL row at bottom\n"
            "7. Click 'Copy New' -> verify creates new version with all data copied\n"
            "8. Click 'New Quotation' -> verify navigates to Quotation Create"
        ),
        "expected": (
            "1. Multi-year view shows Revenue/Cost/Margin for each category across years\n"
            "2. YEAR 2 and YEAR 3 labeled as '(RENEW)' - indicates renewal contracts\n"
            "3. Cell colors: green = positive margin, red = below target\n"
            "4. 'Copy New' creates duplicate P&L for modification\n"
            "5. 'New Quotation' creates quotation directly from approved P&L\n"
            "6. Only latest P&L version is active\n"
            "7. Permissions tab controls: read access, edit access, attachment download access"
        ),
    },
    {
        "id": "TC-017",
        "title": "Quotation Create - Auto-Population from P&L and Opportunity",
        "module": "Quotation",
        "category": "Data Flow",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Approved P&L V1 exists, linked to Opportunity with Customer",
        "steps": (
            "1. Navigate to Opportunity Details -> Quotation(N) tab -> Click 'Create'\n"
            "2. Verify Quotation Create opens as slide-over panel\n"
            "3. Verify Header auto-populated:\n"
            "   - P&L, Customer, Opportunity, Sales Rep, Currency (MYR), Quote Date, Address, Department\n"
            "4. Verify Financial Summary: Total Excl Tax, Service Tax 8%, Grand Total\n"
            "5. Verify Line Items auto-populated from P&L\n"
            "6. Verify T&C template populated with 8 numbered items and placeholders"
        ),
        "expected": (
            "1. Quotation is a 'view' of approved P&L - data flows P&L -> Quotation\n"
            "2. Auto-populated fields: Customer, Opportunity, Currency, Sales Rep, Date, Address, Department\n"
            "3. Financial figures inherited from P&L and recalculated with 8% SST\n"
            "4. Service Tax = Total Excl Tax x 8% (Malaysia SST)\n"
            "5. T&C template contains 8 items with placeholders: <Validity>, <xx>, <Payment Terms>\n"
            "6. Quotation Title autocomplete suggests existing names to prevent duplicates\n"
            "7. P&L lookup is required - cannot create Quotation without approved P&L"
        ),
    },
    {
        "id": "TC-018",
        "title": "Quotation PDF Preview - Document Format and Signature Rendering",
        "module": "Quotation",
        "category": "PDF Export",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Quotation exists with Prepared Signature and Approved Signature uploaded",
        "steps": (
            "1. Navigate to Quotation Details -> Click PDF Preview\n"
            "2. Verify PDF Preview modal displays company header, customer info, line items, financial summary, T&C\n"
            "3. Verify signature blocks: 'Prepare by' and 'Approved by' show uploaded images\n"
            "4. Click 'Export PDF'\n"
            "5. Verify browser download notification with file name and size\n"
            "6. Verify downloaded file name includes Opportunity name"
        ),
        "expected": (
            "1. PDF matches the on-screen preview exactly\n"
            "2. Company header: 'SECUREMETRIC TECHNOLOGY SDN. BHD. (759814-V)'\n"
            "3. Both Prepared by and Approved by signatures rendered in final PDF\n"
            "4. T&C placeholders replaced with actual values at generation time\n"
            "5. Financial summary in PDF matches Quotation financial summary\n"
            "6. Export PDF triggers browser download (not in-app save)\n"
            "7. KNOWN: HTTP environments may block downloads"
        ),
    },
    {
        "id": "TC-019",
        "title": "Quotation Rich Text Editor - T&C and Acceptance Content",
        "module": "Quotation",
        "category": "Rich Text",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Quotation Create form open as slide-over panel",
        "steps": (
            "1. Scroll to Foot section of Quotation Create\n"
            "2. Locate Terms and Conditions rich text editor\n"
            "3. Verify toolbar: Undo, Redo, Copy/Paste, Font Size, Paragraph formatting, Fullscreen\n"
            "4. Verify default template content with 8 numbered items\n"
            "5. Edit T&C content using toolbar\n"
            "6. Locate Acceptance Instruction rich text editor\n"
            "7. Verify 'Prepared by' is read-only: 'Automatically generated by the system'\n"
            "8. Verify 'Approved by' is empty input field"
        ),
        "expected": (
            "1. Rich text editors support full formatting toolbar\n"
            "2. T&C template contains dynamic placeholders replaced at generation time\n"
            "3. Acceptance Instruction can be customized per quotation\n"
            "4. 'Prepared by' is system-generated and cannot be edited\n"
            "5. 'Approved by' remains empty until approval workflow completes\n"
            "6. Both signature uploads are drag-drop zones (jpg/gif/png only)\n"
            "7. Fullscreen mode available for large content editing"
        ),
    },
    {
        "id": "TC-020",
        "title": "Quotation Approval - Circulate Workflow Before Submit",
        "module": "Quotation",
        "category": "Approval",
        "role": "Sales Manager",
        "priority": "P2",
        "precondition": "Quotation exists in draft state",
        "steps": (
            "1. Navigate to Quotation Details -> Process Approvals sidebar\n"
            "2. Set Operation = 'Circulate' (NOT 'submit')\n"
            "3. Set Identity of circulator = 'SMMY' (dropdown)\n"
            "4. Select circulation objects via person icon picker\n"
            "5. Enter Review Comments (verify 0/200 character counter)\n"
            "6. Upload attachment (optional)\n"
            "7. Toggle ON: 'Reading opinions must be responded to'\n"
            "8. Click Submit -> triggers circulation workflow\n"
            "9. Verify circulation objects receive notification\n"
            "10. After circulation, set Operation = 'submit' for final approval"
        ),
        "expected": (
            "1. Circulate workflow allows pre-approval review by designated users\n"
            "2. Circulation objects receive notification to review\n"
            "3. Toggle 'Reading opinions must be responded to' enforces response requirement\n"
            "4. Character counter for Review Comments: 0/200\n"
            "5. Circulation is separate from formal approval\n"
            "6. After circulation completes, standard submit triggers approval workflow\n"
            "7. Process Approvals sidebar supports: Track, Common Comments, Upload, Electronic signature"
        ),
    },
    {
        "id": "TC-021",
        "title": "PO Creation - Mandatory File Upload and Opportunity Status Update",
        "module": "PO",
        "category": "Order Creation",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Quotation approved, Opportunity exists",
        "steps": (
            "1. Navigate to Opportunity Details -> PO(N) tab -> Click 'Create PO'\n"
            "2. Verify auto-populated fields: Opportunity, Customer, Currency, Total Amount (read-only)\n"
            "3. Fill mandatory fields:\n"
            "   - PO Number: customer PO number (e.g., '8596848')\n"
            "   - Date: auto-set to current date\n"
            "   - Quotation: select source quotation\n"
            "   - Deal Category: select (e.g., 'PKI')\n"
            "   - PO File: upload scanned PO document (MANDATORY)\n"
            "4. Verify Product Details table auto-populated from Quotation\n"
            "5. Submit for approval\n"
            "6. Navigate to Opportunity -> verify Status updated to 'Won'\n"
            "7. Verify PO status = 'Active' after approval"
        ),
        "expected": (
            "1. PO creation is a prerequisite for Sales Order creation\n"
            "2. PO File upload is MANDATORY - cannot submit without scanned PO document\n"
            "3. Parent Opportunity status automatically updates to 'Won' upon PO creation\n"
            "4. PO inherits: Customer, Opportunity, Currency, Total Amount, Deal Category\n"
            "5. Product Details table shows items from Quotation line items\n"
            "6. PO is a child record of 'won' Opportunity"
        ),
    },
    {
        "id": "TC-022",
        "title": "SO Payment Schedule - Revenue Split Logic (Software and Hardware % Distribution)",
        "module": "SO",
        "category": "Payment Schedule",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "PO approved, ready to create Sales Order",
        "steps": (
            "1. Navigate to Sales Order Create (from approved Quotation)\n"
            "2. Verify auto-populated fields from Quotation\n"
            "3. Fill mandatory fields: Customer PO No., P.O. Date, Entity\n"
            "4. Scroll to Payment Schedule section\n"
            "5. Set Milestone Type = 'By Products'\n"
            "6. Add 2 payment milestones:\n"
            "   Row 1: Name='Down Payment', Product Info: 'Software 10% Hardware 20%'\n"
            "   Row 2: Name='Final Payment', Product Info: 'Software 90% Hardware 80%'\n"
            "7. Verify Software percentages sum to 100% (10% + 90%)\n"
            "8. Verify Hardware percentages sum to 100% (20% + 80%)\n"
            "9. Verify Receivable Amount auto-calculated: Product Sum x Receivable %\n"
            "10. Submit SO"
        ),
        "expected": (
            "1. Revenue split: each product category percentage sums to 100% across milestones\n"
            "2. Software = 10% (Down) + 90% (Final) = 100%\n"
            "3. Hardware = 20% (Down) + 80% (Final) = 100%\n"
            "4. Receivable Amount = Product Sum x Receivable % for each milestone\n"
            "5. Payment Schedule auto-generates from milestone definitions\n"
            "6. Entity selection auto-populates Entity Code and Department\n"
            "7. Customer PO No. is MANDATORY before submission"
        ),
    },
    {
        "id": "TC-023",
        "title": "SO Details - Statistic Bar and 9-Tab Navigation",
        "module": "SO",
        "category": "Financial Tracking",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Sales Order exists with 2 Payment Schedule milestones",
        "steps": (
            "1. Navigate to SO Details page\n"
            "2. Verify Sales Order Statistic Bar (6 metrics):\n"
            "   Receivable, Received, Uncollected, Invoiced, Uninvoiced, Order Total\n"
            "3. Verify 9 detail tabs with record counts:\n"
            "   Detail Information, Products(6), Delivery(0), Contract(0),\n"
            "   Collection Details(0), Payment Schedule(2), Progress Invoice(0),\n"
            "   Transaction Record(0), Approval Workflow\n"
            "4. Verify Receivable = sum of all Payment Schedule Receivable Amounts\n"
            "5. Verify Uncollected = Receivable - Received\n"
            "6. Verify Uninvoiced = Receivable - Invoiced"
        ),
        "expected": (
            "1. Statistic Bar provides at-a-glance financial status\n"
            "2. All 6 metrics are auto-calculated from underlying records\n"
            "3. Tab record counts reflect actual child records\n"
            "4. Financial calculations verified across all related records\n"
            "5. When no payments/invoices: Receivable = Uncollected = Uninvoiced = Order Total\n"
            "6. Tab navigation allows quick access to related records"
        ),
    },
    {
        "id": "TC-024",
        "title": "Permission - Non-Owner Editing Payment Schedule (Known Bug Regression)",
        "module": "SO",
        "category": "Permission Bug",
        "role": "Sales Manager",
        "priority": "P1",
        "precondition": "SO exists, Payment Schedule created, User 'Yuwin' is NOT person in charge",
        "steps": (
            "1. Login as 'Yuwin' (not the P.I.C. of this SO)\n"
            "2. Navigate to SO Details -> Payment Schedule tab\n"
            "3. Click on a Payment Schedule milestone\n"
            "4. Attempt to edit Payment Schedule details\n"
            "5. Click 'Change Completion %' -> attempt to modify\n"
            "6. Login as P.I.C. -> attempt same actions\n"
            "7. Verify which user can and cannot edit"
        ),
        "expected": (
            "1. BUG: Non-owner (Yuwin) CAN edit Payment Schedule despite not being P.I.C.\n"
            "   - Expected: Only person in charge can edit\n"
            "   - Actual: Any user with SO access can edit Payment Schedule\n"
            "2. Bug severity: HIGH - financial data integrity at risk\n"
            "3. P.I.C. can edit (expected behavior)\n"
            "4. Payment Schedule contains critical financial data (Receivable amounts, Completion %)\n"
            "5. Contradicts Service Team permission model described in video 16"
        ),
    },
    {
        "id": "TC-025",
        "title": "Contract - Multiple Expiry Notification Rules",
        "module": "Contract",
        "category": "Expiry Reminder",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "SO exists, ready to create Contract",
        "steps": (
            "1. Navigate to SO Details -> Contract tab -> Click 'Create'\n"
            "2. Fill Contract form:\n"
            "   - Contract Title: 'ABC NDA'\n"
            "   - Sales Order: auto-filled\n"
            "   - Customer: auto-filled\n"
            "   - Signed Date: '04/29/2026', Expiry Date: '04/28/2027'\n"
            "3. Configure 3 reminder rules:\n"
            "   - Rule 1: Notify 'Affendi', 30 days before expiry\n"
            "   - Rule 2: Notify 'CK', 10 days before expiry\n"
            "   - Rule 3: Notify 'Affendi' and 'CK', 5 days before expiry\n"
            "4. Upload attachment (optional)\n"
            "5. Save Contract\n"
            "6. Verify Contract appears in Contract tab with Serial No. and Title"
        ),
        "expected": (
            "1. Multiple reminder rules can be configured per contract\n"
            "2. Each rule has: Notify Who (user lookup) + Days Before Expiry\n"
            "3. Notifications trigger at configured intervals before expiry date\n"
            "4. Contract ID is auto-generated on save\n"
            "5. Contract file syncs to Customer 360\n"
            "6. Contract is created from SO - Sales Order and Customer pre-filled"
        ),
    },
    {
        "id": "TC-026",
        "title": "Payment Schedule - Receivable Calculation and Completion % Tracking",
        "module": "Payment Schedule",
        "category": "Financial Calculation",
        "role": "PM",
        "priority": "P1",
        "precondition": "SO exists with 2 Payment Schedule milestones (Down 50%, Final 50%)",
        "steps": (
            "1. Navigate to SO Details -> Payment Schedule tab -> click a milestone\n"
            "2. Verify Payment Schedule Details drawer opens\n"
            "3. Verify financial fields:\n"
            "   - ID: PP+YYYYMMDD+seq (e.g., PP202604270002)\n"
            "   - Receivable amount = Order Total x Receivable % (37,709.74 x 50% = 18,854.87)\n"
            "   - Invoiced Amount: 0.00, Uninvoiced = Receivable - Invoiced\n"
            "   - Amount Uncollected = Receivable - Received\n"
            "4. Verify Status = 'Uncompleted' (blue badge)\n"
            "5. Verify Payment Status = 'Unpaid' (blue badge)\n"
            "6. Click 'Change Completion %' -> set to 100%\n"
            "7. Verify system alerts AM to generate Proforma Invoice\n"
            "8. Verify tabs: Detail Information | Payment Detail(0) | Billing Detail(0) | System Record"
        ),
        "expected": (
            "1. Financial calculation:\n"
            "   - Receivable amount = Order Total x Receivable %\n"
            "   - Uninvoiced Amount = Receivable - Invoiced Amount\n"
            "   - Amount Uncollected = Receivable - Received Amount\n"
            "2. At 100% Completion %, system prompts AM to generate Invoice Application\n"
            "3. Status transitions: 'Uncompleted' -> (received payment) -> 'Completed'\n"
            "4. BUG: 'Contract' field label should be 'Sales Order'\n"
            "5. PM updates Completion %; at 100%, system alerts AM to generate PI"
        ),
    },
    {
        "id": "TC-027",
        "title": "Invoice Application - Project vs Milestone Billing Comparison",
        "module": "Invoice Application",
        "category": "Billing Types",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "SO exists with 2 Payment Schedule milestones, PM confirmed 100% completion",
        "steps": (
            "1. Navigate to Invoice Application Create\n"
            "2. Select Invoice Type = 'Project'\n"
            "3. Select Sales Order -> verify auto-populated: Customer, Currency, P.I.C., Opportunity\n"
            "4. Verify Line Items auto-populated from SO (Service Period, Product, Qty, Unit Price)\n"
            "5. Verify Invoice Amount: Total Excl Tax + Service Tax @ 8% = Grand Total\n"
            "6. Save -> Click PDF Preview -> verify Proforma Invoice format\n"
            "7. Delete this IR -> Create new with Invoice Type = 'Milestone'\n"
            "8. Verify Milestone section: Payment Schedule ID 'Select record' modal\n"
            "9. In modal: verify filter fields and grid columns\n"
            "10. Select Payment Schedule -> verify milestone data auto-populates"
        ),
        "expected": (
            "1. Project Type: Full-value invoicing, line items auto-filled from SO\n"
            "2. Milestone Type: Progress billing, requires selecting from Payment Schedule records\n"
            "3. Select Record Modal filters by: SO ID, Payment Type, Receivable amount, PS ID\n"
            "4. Grid columns: Serial No., Payment Schedule ID, Milestone Name, Job Content\n"
            "5. Both types auto-populate: Invoice Application ID, PI No., Ref No.\n"
            "6. Both types apply 8% SST automatically\n"
            "7. UI Typo: 'Univoiced Amount' should be 'Uninvoiced Amount'"
        ),
    },
    {
        "id": "TC-028",
        "title": "Invoice Application - PDF Preview and Digital Signature",
        "module": "Invoice Application",
        "category": "PDF Export",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Invoice Application exists, in approval workflow",
        "steps": (
            "1. Navigate to Invoice Application Details\n"
            "2. Click PDF Preview\n"
            "3. Verify Proforma Invoice format:\n"
            "   - Company logo: 'SECURE METRIC TECHNOLOGY'\n"
            "   - Recipient Block, Invoice Metadata, Line Items Table\n"
            "4. Click Export PDF -> verify download\n"
            "5. Navigate to Process Approvals sidebar\n"
            "6. Verify Digital Signature Pad (canvas-based)\n"
            "7. Draw signature -> verify Preview | Delete options\n"
            "8. Submit for approval -> verify BPM workflow triggered"
        ),
        "expected": (
            "1. PDF Preview shows Proforma Invoice format (different from Quotation PDF)\n"
            "2. Line Items include service period (Year1, etc.)\n"
            "3. Digital Signature Pad uses canvas-based drawing (not file upload)\n"
            "4. Signature Preview shows drawn signature before submission\n"
            "5. Approval workflow same pattern as other modules\n"
            "6. IR ID format: IRYYYYMMDDNNNN (e.g., IR202605270004)\n"
            "7. Both PDF Preview and Export PDF available in the modal"
        ),
    },
    {
        "id": "TC-029",
        "title": "Collection - Finance Confirmation and Bank Receipt",
        "module": "Collection",
        "category": "Finance Workflow",
        "role": "Finance User",
        "priority": "P2",
        "precondition": "Invoice Application approved, customer payment received",
        "steps": (
            "1. Navigate to Collection Create\n"
            "2. Fill required fields:\n"
            "   - Customer Name: select customer account\n"
            "   - Payment Date: select payment date\n"
            "   - Collection Amount (Total): enter amount\n"
            "   - Settlement Currency: select currency (CFG dropdown)\n"
            "3. Fill optional fields: Owner, Department, Finance Confirmer, Bank Receipt\n"
            "4. Submit for finance confirmation\n"
            "5. Login as Finance Confirmer -> verify collection in pending list\n"
            "6. Confirm collection -> verify fd_finance_confirm_time recorded\n"
            "7. Verify Collection Amount matches Invoice Application Grand Total"
        ),
        "expected": (
            "1. Collection requires finance user confirmation workflow\n"
            "2. Bank receipt attachment is optional but recommended for audit trail\n"
            "3. Currency selection uses CFG dropdown (system-configured currencies)\n"
            "4. Finance confirmation records: fd_finance_employee_id + fd_finance_confirm_time\n"
            "5. Collection Amount should match the Invoice Application amount being collected\n"
            "6. Lock Status can be set after confirmation to prevent further edits\n"
            "7. Chinese labels in UI are defects (per Wiki language policy)"
        ),
    },
    {
        "id": "TC-030",
        "title": "Delivery - SO-Linked Shipment with Courier Tracking",
        "module": "Delivery",
        "category": "SO Linkage",
        "role": "Operations",
        "priority": "P2",
        "precondition": "SO exists, products ready for delivery",
        "steps": (
            "1. Navigate to Delivery Create\n"
            "2. Fill required field: Sales Order -> select SO\n"
            "3. Verify auto-populated: Customer (from SO), Sales Person (from SO)\n"
            "4. Fill delivery details: Date, PI No., Self Collect ID, Item Description, Ship Date\n"
            "5. Fill courier info: Courier Services, Air Way Bill, Weight\n"
            "6. Fill shipping costs: Fuel Surcharge (MYR), Shipping Cost (MYR)\n"
            "7. Submit delivery record"
        ),
        "expected": (
            "1. Delivery is linked to SO - Customer and Sales Person auto-populated\n"
            "2. Product items restricted to SO products only\n"
            "3. Courier tracking: Courier Services + Air Way Bill + Weight\n"
            "4. Shipping costs tracked in MYR (Fuel Surcharge + Shipping Cost)\n"
            "5. Delivery managed by operations team via dedicated view\n"
            "6. PI should be created before Delivery (per audio workflow)\n"
            "7. Chinese labels in UI are defects"
        ),
    },
    {
        "id": "TC-031",
        "title": "Task - Complete Owner/Executor/CC Workflow with Feedback Loop",
        "module": "Task",
        "category": "Full Lifecycle",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "Opportunity exists with Owner 'Affendi' and executor 'Danny'",
        "steps": (
            "1. Navigate to Opportunity Details -> Click 'New Task'\n"
            "2. Fill Task Create form:\n"
            "   - Deadline: '06/06/2026 11:00 pm', Owner: auto-filled 'Affendi'\n"
            "   - Executor: 'Danny', Priority: 'High'\n"
            "   - Associated Type: auto-filled 'Opportunity'\n"
            "   - Description: 'Prepare solution proposal for CIMB eKYC'\n"
            "   - C.C. Recipient: select 'CK'\n"
            "3. Configure Reminder: '1 day before', Method = To-Do + Email\n"
            "4. Save Task -> verify task appears in Opportunity Task list\n"
            "5. Login as Danny (Executor) -> check To-Do list\n"
            "6. Click 'Task Feedback' -> set Completion Progress to 60%\n"
            "7. Enter feedback -> Check Owner (notify) -> Confirm\n"
            "8. Login as Affendi (Owner) -> verify status = 'To be confirm'\n"
            "9. Review feedback -> Click 'Pass'\n"
            "10. Verify status = 'Completed', Progress = 100%"
        ),
        "expected": (
            "1. Task created with auto-filled Associated Type and Associated Data\n"
            "2. Executor receives To-Do item AND email notification\n"
            "3. C.C. Recipient receives notification but cannot edit or confirm\n"
            "4. Task Feedback: Completion Progress slider (0-100%) with rich text\n"
            "5. Status transitions: In progress -> To be confirm -> Completed\n"
            "6. Owner 'Pass' action marks task as completed (100% progress)\n"
            "7. 'Task completed' button is disabled after completion\n"
            "8. Task Feedback tab shows timeline of all feedback entries"
        ),
    },
    {
        "id": "TC-032",
        "title": "Task - Sub-Task Creation and Hierarchy",
        "module": "Task",
        "category": "Sub-Task",
        "role": "Sales Rep",
        "priority": "P3",
        "precondition": "Task exists in 'In progress' status",
        "steps": (
            "1. Navigate to Task Details -> Basic Info tab\n"
            "2. Locate 'Sub-tasks' section -> Click '+ Add Sub-tasks'\n"
            "3. Create 2 sub-tasks:\n"
            "   - Sub-task 1: 'Research competitor solutions', Deadline: 2 days, Priority: Medium\n"
            "   - Sub-task 2: 'Draft proposal outline', Deadline: 3 days, Priority: High\n"
            "4. Verify sub-tasks appear under parent task\n"
            "5. Verify parent task shows 'Superior Task: -' (no parent)\n"
            "6. Verify sub-tasks show 'Superior Task: [Parent Task Name]'\n"
            "7. Complete sub-task 1 -> verify parent task progress updates\n"
            "8. Complete sub-task 2 -> verify parent task can be marked completed"
        ),
        "expected": (
            "1. Sub-tasks follow same workflow as parent tasks (Owner/Executor/Feedback)\n"
            "2. Superior Task field links child to parent\n"
            "3. Parent task progress may be aggregate of sub-task progress (needs verification)\n"
            "4. Sub-tasks can have their own Executors (different from parent)\n"
            "5. Sub-tasks can have their own reminders and C.C. recipients\n"
            "6. Parent task cannot be completed until all sub-tasks are completed (needs verification)"
        ),
    },
    {
        "id": "TC-033",
        "title": "Task - Multiple Reminder Rules with Dual-Channel Notification",
        "module": "Task",
        "category": "Reminder",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Task exists with deadline in 3 days",
        "steps": (
            "1. Navigate to Task Create or Edit\n"
            "2. Locate Reminder sub-table\n"
            "3. Add 3 reminder rules:\n"
            "   Row 1: '3 hours before' -> Method: To-Do + Email\n"
            "   Row 2: '1 hour before' -> Method: To-Do only\n"
            "   Row 3: 'At deadline' -> Method: Email only\n"
            "4. Save task\n"
            "5. Verify reminder table shows all 3 rules\n"
            "6. Wait for reminder triggers (or simulate)\n"
            "7. Verify To-Do notifications at 3 hours and 1 hour before\n"
            "8. Verify email notifications at 3 hours before and at deadline"
        ),
        "expected": (
            "1. Multiple reminder rules can be added per task\n"
            "2. Reminder timing calculated relative to Deadline\n"
            "3. Dual-channel: To-Do creates item in task list, Email sends notification\n"
            "4. Methods can be mixed per rule (To-Do only, Email only, or both)\n"
            "5. Reminder options: 1 day / 3 hours / 1 hour / 30 min / 15 min / At deadline\n"
            "6. Row actions: Insert (add above) | Copy (duplicate rule)"
        ),
    },
    {
        "id": "TC-034",
        "title": "Pipeline Kanban - 6-Stage Summary Bar and Filter Interaction",
        "module": "Pipeline Kanban",
        "category": "Dashboard",
        "role": "Sales Manager",
        "priority": "P2",
        "precondition": "System has deals across all 6 pipeline stages",
        "steps": (
            "1. Navigate to CRM -> Pipeline\n"
            "2. Verify 6-stage summary bar: LEAD, OPPORTUNITY, QUOTATION, PO, SALES ORDER, PAYMENT\n"
            "3. Verify each tile shows: Deal count + Pipeline value (MYR)\n"
            "4. Click 'QUOTATION' tile -> verify list filters to QUOTATION stage only\n"
            "5. Verify 'Reset' button restores full view\n"
            "6. Test TIME filter: 'This Month' vs 'This Year' vs Custom\n"
            "7. Test ENTITY filter: filter by SCMY vs SMMY\n"
            "8. Test REP filter: filter by specific sales rep\n"
            "9. Verify summary line updates: 'Showing N deals, Pipeline MYR X,XXX,XXX.XX'\n"
            "10. Click any project row -> verify navigates to detail page"
        ),
        "expected": (
            "1. Pipeline Kanban shows 6 stages: LEAD -> OPPORTUNITY -> QUOTATION -> PO -> SALES ORDER -> PAYMENT\n"
            "2. Stage tiles filter the project list below\n"
            "3. All monetary amounts in branch default currency (MYR)\n"
            "4. TIME filter supports: This Year / This Month / Custom date range\n"
            "5. ENTITY filter scopes to specific legal entity\n"
            "6. REP filter scopes to specific sales rep\n"
            "7. Dot-track progress indicator: filled = completed stages, ring = current stage\n"
            "8. DOCX says 5 stages but actual UI shows 6 (QUOTATION is separate)"
        ),
    },
    {
        "id": "TC-035",
        "title": "Duplicate Check - Fuzzy Search and Yellow Highlight Matching",
        "module": "Duplicate Check",
        "category": "Cross-Entity Search",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "System has existing customers and contacts with similar names",
        "steps": (
            "1. Navigate to CRM -> DUPLICATE CHECK\n"
            "2. Search for partial name (e.g., 'Elite' when 'Elite Enterprise Solutions' exists)\n"
            "3. Verify results summary bar: Total Records Found, Contacts(N), Customers(N)\n"
            "4. Verify split tabs: Customers(N) and Contacts(N) with count badges\n"
            "5. Verify matching text highlighted in yellow\n"
            "6. Verify result columns: No., Name, Owner, Entity, Legal ID\n"
            "7. Switch between Customers and Contacts tabs\n"
            "8. Search for email address -> verify cross-entity match\n"
            "9. Search for Registration Code -> verify exact match\n"
            "10. Verify pagination"
        ),
        "expected": (
            "1. Single search queries BOTH Contacts and Customers simultaneously\n"
            "2. Fuzzy matching returns partial matches (not exact only)\n"
            "3. Matching text highlighted in yellow in results\n"
            "4. Results split by entity type with count badges\n"
            "5. Returns only real database records (not drafts)\n"
            "6. Search targets: Name, Mobile, Email, Registration Code\n"
            "7. Duplicate Check should be used BEFORE creating new records"
        ),
    },
    {
        "id": "TC-036",
        "title": "End-to-End Pipeline - Lead to Payment Complete Data Flow",
        "module": "Cross-Module",
        "category": "Full Pipeline",
        "role": "Sales Rep",
        "priority": "P0",
        "precondition": "Fresh test environment with no existing records",
        "steps": (
            "1. CREATE Lead -> Assign to Sales Rep -> Fill Principal Allocation\n"
            "2. CONVERT Lead -> Verify Customer + Opportunity created\n"
            "3. Verify carry-over: Service Team, Activity Logs, Contacts\n"
            "4. Navigate to Opportunity -> Advance stage to 'Proposal/POC' (50%)\n"
            "5. CREATE P&L from Opportunity -> Add products -> Verify Financial Summary\n"
            "6. SUBMIT P&L -> Verify approval routing (margin-based)\n"
            "7. APPROVE P&L -> CREATE Quotation -> Verify auto-population from P&L\n"
            "8. Export Quotation PDF -> Verify signatures and financial data\n"
            "9. CREATE PO -> Upload PO file -> Verify Opportunity status = 'Won'\n"
            "10. CREATE SO -> Define Payment Schedule (Down 50% + Final 50%)\n"
            "11. CREATE Contract -> Configure expiry reminders\n"
            "12. PM updates Payment Schedule Completion % to 100%\n"
            "13. CREATE Invoice Application (Milestone) -> Select Payment Schedule\n"
            "14. APPROVE Invoice Application -> Record Collection\n"
            "15. CREATE Delivery -> Link to SO -> Enter courier info"
        ),
        "expected": (
            "1. Complete data lineage: Lead -> Customer -> Contact -> Opportunity -> P&L -> Quotation -> PO -> SO -> Contract -> Payment Schedule -> Invoice Application -> Collection -> Delivery\n"
            "2. At each stage, verify auto-populated fields match parent record\n"
            "3. Financial amounts consistent across modules (no rounding errors)\n"
            "4. Status transitions correct at each stage\n"
            "5. Service Team and permissions enforced throughout pipeline\n"
            "6. Pipeline Kanban shows deal progressing through 6 stages\n"
            "7. Total pipeline time from Lead creation to Payment receipt"
        ),
    },
    {
        "id": "TC-037",
        "title": "Currency Consistency - MYR/IDR Dual-Track Across Pipeline",
        "module": "Cross-Module",
        "category": "Currency Consistency",
        "role": "Sales Rep",
        "priority": "P2",
        "precondition": "Lead created with Currency=IDR, Exchange Rate=0.000223",
        "steps": (
            "1. Create Lead with Currency=IDR, Estimated Amount=50,000,000\n"
            "2. Convert to Opportunity -> verify Currency=IDR\n"
            "3. Create P&L -> verify Currency=IDR inherited\n"
            "4. Add products with MYR prices -> verify conversion to IDR\n"
            "5. Create Quotation -> verify Currency=IDR, prices in IDR\n"
            "6. Create SO -> verify Currency=IDR, Order Amount in IDR\n"
            "7. Create Invoice Application -> verify Currency=IDR\n"
            "8. Create Collection -> verify Settlement Currency=IDR\n"
            "9. Navigate to Pipeline Kanban -> verify amount displayed in MYR"
        ),
        "expected": (
            "1. Currency inherited at each stage from parent record\n"
            "2. MYR is the unified reporting currency (Blueprint Design Principle 5)\n"
            "3. Pipeline Kanban displays all amounts in branch default currency (MYR)\n"
            "4. Exchange rate conversion accurate at each step\n"
            "5. Collection can record payment in different currency (settlement currency)\n"
            "6. Currency switching in P&L recalculates ALL figures"
        ),
    },
    {
        "id": "TC-038",
        "title": "Data Sharding - Cross-Entity Visibility and Public Pool",
        "module": "Cross-Module",
        "category": "Entity Isolation",
        "role": "Sales Manager",
        "priority": "P1",
        "precondition": "Customer exists in SCMY entity, another in SMMY entity",
        "steps": (
            "1. Login as User-A (SCMY entity)\n"
            "2. Navigate to Customer List -> verify only SCMY customers visible\n"
            "3. Attempt to access SMMY customer via direct URL -> access denied\n"
            "4. Navigate to Public Pool -> verify unassigned SCMY customers visible\n"
            "5. Verify SMMY Public Pool customers NOT visible to SCMY user\n"
            "6. Login as User-B (SMMY entity)\n"
            "7. Verify mirror behavior: SMMY customers visible, SCMY not visible\n"
            "8. Admin user: verify can see ALL entities records"
        ),
        "expected": (
            "1. Entity isolation: each entity data is mutually invisible (Blueprint Principle 3)\n"
            "2. Public Pool visible to all AMs in the SAME region only\n"
            "3. Direct URL access to cross-entity records returns access denied\n"
            "4. Admin/super-admin can see all entities\n"
            "5. Export restrictions protect sensitive contact data per entity\n"
            "6. Multi-country sharding: MY/VN/PH/ID data partitioned"
        ),
    },
    {
        "id": "TC-039",
        "title": "P&L Version Locking - Quote Generation and Version Correspondence",
        "module": "Cross-Module",
        "category": "P&L Version Control",
        "role": "Sales Rep",
        "priority": "P1",
        "precondition": "P&L V1 approved, Quotation V1 generated",
        "steps": (
            "1. Navigate to P&L Details -> verify Version = 'V1'\n"
            "2. Navigate to Quotation Details -> verify linked to P&L V1\n"
            "3. Click 'Copy New' on P&L -> creates P&L V2\n"
            "4. Modify pricing in P&L V2 -> Submit -> Approve\n"
            "5. Navigate to Quotation -> verify still linked to P&L V1\n"
            "6. Click 'New Quotation' from P&L V2 -> creates Quotation V2\n"
            "7. Verify Quotation V2 reflects P&L V2 pricing\n"
            "8. Verify Quotation V1 still shows P&L V1 pricing (version-locked)\n"
            "9. Attempt to edit Quotation V1 directly -> should be blocked\n"
            "10. Verify only 'active' P&L version (V2) included in Pipeline Kanban reports"
        ),
        "expected": (
            "1. P&L V1 -> Quotation V1, P&L V2 -> Quotation V2 (version correspondence)\n"
            "2. Quotations cannot be edited directly - must modify underlying P&L\n"
            "3. 'Copy New' creates new P&L version with all data copied\n"
            "4. Historical P&L versions are deactivated\n"
            "5. Only 'active' P&L version data aggregated in Pipeline Kanban\n"
            "6. Version locking prevents data inflation in pipeline reports\n"
            "7. Full audit trail: every version change tracked in Process Record"
        ),
    },
]

# ── workbook builder ──────────────────────────────────────────────────────────
def build_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.sheet_view.showGridLines = True

    COLS = ["TC ID", "Title", "Module", "Category", "Role", "Priority",
            "Precondition", "Test Steps", "Expected Results",
            "Acceptance Status", "Remarks"]
    COL_WIDTHS = [10, 52, 16, 22, 14, 10, 40, 60, 60, 18, 30]

    # ── header row ─────────────────────────────────────────────────────────
    for col_idx, (label, width) in enumerate(zip(COLS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = Font(bold=True, color=HEADER_FG, size=11, name="Calibri")
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = thin_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28

    # ── data rows ──────────────────────────────────────────────────────────
    for row_idx, tc in enumerate(TEST_CASES, start=2):
        bg_hex = ALT_BG if row_idx % 2 == 0 else "FFFFFF"

        values = [
            tc["id"], tc["title"], tc["module"], tc["category"],
            tc["role"], tc["priority"], tc["precondition"],
            tc["steps"], tc["expected"],
            "",   # Acceptance Status - to be filled manually
            "",   # Remarks - to be filled manually
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border()
            cell.font = Font(size=10, name="Calibri")

            # Priority cell gets coloured badge
            if col_idx == 6:
                fg, txt = pri_colour(val)
                cell.fill = PatternFill("solid", fgColor=fg)
                cell.font = Font(bold=True, size=10, color=txt, name="Calibri")
                cell.alignment = Alignment(horizontal="center",
                                           vertical="center")
            else:
                cell.fill = PatternFill("solid", fgColor=bg_hex)

        # row height - taller for steps/expected
        ws.row_dimensions[row_idx].height = 180

    # ── Acceptance Status dropdown (column J = col 10) ─────────────────────
    last_row = 1 + len(TEST_CASES)
    dv = DataValidation(
        type="list",
        formula1='"Pass,Fail,Blocked,Not Tested"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.sqref = f"J2:J{last_row}"
    ws.add_data_validation(dv)

    # ── freeze panes & auto-filter ─────────────────────────────────────────
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"

    # ── summary sheet ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2.sheet_view.showGridLines = False

    # Title
    ws2["A1"] = "Securemetric CRM - Test Case Summary"
    ws2["A1"].font = Font(bold=True, size=16, color=HEADER_BG, name="Calibri")
    ws2["A1"].alignment = Alignment(horizontal="left")
    ws2.row_dimensions[1].height = 30
    ws2.merge_cells("A1:D1")

    ws2["A2"] = "Generated: 2026-06-05"
    ws2["A2"].font = Font(italic=True, size=10, color="666666", name="Calibri")

    # Priority summary
    headers = ["Priority", "Count", "Description"]
    summary_data = [
        ["P0", 1, "Critical - must pass before any release"],
        ["P1", 20, "High - core business logic, bug regression, permission enforcement"],
        ["P2", 17, "Medium - secondary features, UI/UX, edge cases"],
        ["P3", 1, "Low - exploratory, needs further clarification"],
        ["TOTAL", 39, ""],
    ]

    for col_i, h in enumerate(headers, start=1):
        c = ws2.cell(row=4, column=col_i, value=h)
        c.font = Font(bold=True, color=HEADER_FG, name="Calibri")
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()

    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 60

    for r_i, row in enumerate(summary_data, start=5):
        pri = row[0]
        for c_i, val in enumerate(row, start=1):
            cell = ws2.cell(row=r_i, column=c_i, value=val)
            cell.border = thin_border()
            cell.alignment = Alignment(horizontal="center" if c_i < 3 else "left",
                                       vertical="center")
            cell.font = Font(size=10, name="Calibri")
            if c_i == 1 and pri != "TOTAL":
                fg, txt = pri_colour(pri)
                cell.fill = PatternFill("solid", fgColor=fg)
                cell.font = Font(bold=True, size=10, color=txt, name="Calibri")
            elif pri == "TOTAL":
                cell.fill = PatternFill("solid", fgColor=HEADER_BG)
                cell.font = Font(bold=True, size=10, color=HEADER_FG, name="Calibri")

    # Module coverage
    module_data = [
        ["P&L", 6, "15.4%"],
        ["Lead", 5, "12.8%"],
        ["Opportunity", 5, "12.8%"],
        ["Quotation", 4, "10.3%"],
        ["Cross-Module", 4, "10.3%"],
        ["SO", 3, "7.7%"],
        ["Task", 3, "7.7%"],
        ["Invoice Application", 2, "5.1%"],
        ["PO", 1, "2.6%"],
        ["Contract", 1, "2.6%"],
        ["Payment Schedule", 1, "2.6%"],
        ["Collection", 1, "2.6%"],
        ["Delivery", 1, "2.6%"],
        ["Pipeline Kanban", 1, "2.6%"],
        ["Duplicate Check", 1, "2.6%"],
    ]

    ws2["A11"] = "Coverage by Module"
    ws2["A11"].font = Font(bold=True, size=12, color=HEADER_BG, name="Calibri")

    mod_headers = ["Module", "Test Cases", "Percentage"]
    for c_i, h in enumerate(mod_headers, start=1):
        c = ws2.cell(row=12, column=c_i, value=h)
        c.font = Font(bold=True, color=HEADER_FG, name="Calibri")
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()

    for r_i, row in enumerate(module_data, start=13):
        for c_i, val in enumerate(row, start=1):
            cell = ws2.cell(row=r_i, column=c_i, value=val)
            cell.border = thin_border()
            bg = ALT_BG if (r_i % 2 == 0) else "FFFFFF"
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.alignment = Alignment(horizontal="center" if c_i > 1 else "left",
                                       vertical="center")
            cell.font = Font(size=10, name="Calibri")

    return wb


if __name__ == "__main__":
    out_path = "CRM_TestCases_FullSystem_EN_2026-06-05.xlsx"
    wb = build_workbook()
    wb.save(out_path)
    print(f"✓ Saved: {out_path}")
    print(f"  {len(TEST_CASES)} test cases | 2 sheets (Test Cases + Summary)")
