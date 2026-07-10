#!/usr/bin/env python3
"""
Generate Excel test cases from JSON data.

Usage:
    cd wiki/test-cases/excel
    python3 generate_excel.py

Output (to ../en/excel/ and ../zh/excel/):
    test-cases-zh-v2.xlsx   — Chinese version v2
    test-cases-en-v2.xlsx   — English version v2
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import os
from pathlib import Path

# All paths are relative to this script's directory
BASE_DIR = Path(__file__).resolve().parent
EN_DIR = BASE_DIR.parent / 'en' / 'excel'
ZH_DIR = BASE_DIR.parent / 'zh' / 'excel'


def create_excel(test_cases: list, output_path: Path, title: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title

    # ── Styles ──────────────────────────────────────────────────
    header_font      = Font(name='Arial', bold=True, size=12, color='FFFFFF')
    header_fill      = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    cell_font      = Font(name='Arial', size=11)
    cell_alignment = Alignment(vertical='top', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'),  bottom=Side(style='thin'),
    )

    # ── Headers ─────────────────────────────────────────────────
    is_en = 'English' in title or title.startswith('Test')
    headers = (
        ['Test ID', 'Module', 'Test Name', 'Priority', 'Precondition',
         'Test Steps', 'Expected Result', 'Status', 'Notes']
        if is_en else
        ['测试编号', '模块', '测试名称', '优先级', '前置条件',
         '测试步骤', '预期结果', '状态', '备注']
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_alignment
        cell.border    = thin_border

    # ── Data rows ────────────────────────────────────────────────
    for row_idx, tc in enumerate(test_cases, 2):
        row_data = [
            tc['id'], tc['module'], tc['name'], tc['priority'],
            tc['precondition'], tc['steps'], tc['expected'],
            tc['status'], tc['notes'],
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value))
            cell.font      = cell_font
            cell.alignment = cell_alignment
            cell.border    = thin_border

            # Priority colour coding
            if col_idx == 4:
                if value == 'P0':
                    cell.fill = PatternFill(start_color='FF6B6B', end_color='FF6B6B', fill_type='solid')
                elif value == 'P1':
                    cell.fill = PatternFill(start_color='FFD93D', end_color='FFD93D', fill_type='solid')

    # ── Layout ──────────────────────────────────────────────────
    for i, width in enumerate([12, 18, 30, 10, 25, 50, 40, 12, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.row_dimensions[1].height = 30
    for row in range(2, len(test_cases) + 2):
        ws.row_dimensions[row].height = 60

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:I{len(test_cases) + 1}'

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"✅ {output_path.name}  ({len(test_cases)} cases)")


def load_json(name: str) -> list:
    path = BASE_DIR / name
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def main():
    jobs = [
        (ZH_DIR / 'test-cases-zh-v2.json', ZH_DIR / 'test-cases-zh-v2.xlsx', '测试用例 v2 (中文)'),
        (EN_DIR / 'test-cases-en-v2.json', EN_DIR / 'test-cases-en-v2.xlsx', 'Test Cases v2 (English)'),
    ]

    for json_path, xlsx_path, title in jobs:
        if not json_path.exists():
            print(f"⚠️  Skipped (not found): {json_path}")
            continue
        data = load_json(json_path)
        create_excel(data, xlsx_path, title)

    print("\n🎉 Done!")


if __name__ == '__main__':
    main()
